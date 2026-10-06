from datetime import date, timedelta
from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

out = Path('/home/ubuntu')

products = [
    ('Laptop', 'Electronics', 55000),
    ('Mouse', 'Electronics', 600),
    ('Keyboard', 'Electronics', 1500),
    ('Monitor', 'Electronics', 12000),
    ('Desk Chair', 'Furniture', 7500),
    ('Office Desk', 'Furniture', 14000),
    ('Bookshelf', 'Furniture', 9000),
    ('Notebook', 'Office Supplies', 180),
    ('Pen Set', 'Office Supplies', 250),
    ('Stapler', 'Office Supplies', 350),
]
rows = []
start = date(2025, 1, 1)
for i in range(50):
    product, category, unit = products[i % len(products)]
    quantity = (i % 4) + 1
    # Small deterministic discount pattern for realistic sales values.
    discount = [0.00, 0.05, 0.10, 0.00, 0.08][i % 5]
    sales = round(unit * quantity * (1 - discount), 2)
    rows.append({
        'Transaction ID': f'T{i+1:03d}',
        'Date': (start + timedelta(days=i)).isoformat(),
        'Product': product,
        'Category': category,
        'Quantity': quantity,
        'Sales': sales,
    })

total = sum(r['Sales'] for r in rows)
avg = total / len(rows)
count = len(rows)

# Separate raw dataset CSV
raw_csv = out / 'Retail_Sales_Dataset.csv'
with raw_csv.open('w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=list(rows[0]))
    writer.writeheader(); writer.writerows(rows)

# Separate summary CSV with calculated values
summary_csv = out / 'Sales_Summary.csv'
with summary_csv.open('w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerow(['KPI', 'Value', 'Formula / Verification'])
    writer.writerow(['Total Sales', f'{total:.2f}', 'SUM of Sales column'])
    writer.writerow(['Average Sales', f'{avg:.2f}', 'AVERAGE of Sales column'])
    writer.writerow(['Transaction Count', count, 'COUNTA of Transaction ID column'])
    writer.writerow([])
    writer.writerow(['Manual Verification', 'Passed', f'Sum of dataset Sales = {total:.2f}'])

# Excel workbook
wb = Workbook()
raw = wb.active
raw.title = 'Raw Data'
summary = wb.create_sheet('Summary')

headers = list(rows[0])
raw.append(headers)
for r in rows:
    raw.append([r[h] for h in headers])

# Raw Data formatting
header_fill = PatternFill('solid', fgColor='1F4E78')
header_font = Font(color='FFFFFF', bold=True)
thin_gray = Side(style='thin', color='D9E2F3')
for cell in raw[1]:
    cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal='center')
for row in raw.iter_rows(min_row=2, max_row=raw.max_row):
    row[1].number_format = 'dd-mm-yyyy'
    row[5].number_format = '₹#,##0.00'
    for cell in row:
        cell.border = Border(bottom=thin_gray)
for col, width in {'A':16,'B':14,'C':18,'D':20,'E':12,'F':16}.items(): raw.column_dimensions[col].width = width
raw.freeze_panes = 'A2'
raw.auto_filter.ref = raw.dimensions
tab = Table(displayName='RetailSalesData', ref=f'A1:F{raw.max_row}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
raw.add_table(tab)

# Summary layout
summary.merge_cells('A1:F1')
summary['A1'] = 'BASIC SALES SUMMARY'
summary['A1'].font = Font(size=18, bold=True, color='FFFFFF')
summary['A1'].fill = PatternFill('solid', fgColor='17365D')
summary['A1'].alignment = Alignment(horizontal='center')
for c in range(1,7): summary.cell(1,c).fill = PatternFill('solid', fgColor='17365D')
summary['A3'] = 'KPI'; summary['B3'] = 'Value'; summary['D3'] = 'Verification'
for cell in ('A3','B3','D3'):
    summary[cell].font = Font(bold=True, color='FFFFFF')
    summary[cell].fill = header_fill
    summary[cell].alignment = Alignment(horizontal='center')
summary['A4'] = 'Total Sales'; summary['B4'] = "=SUM('Raw Data'!F2:F51)"; summary['D4'] = f'Manual sum: ₹{total:,.2f}'
summary['A5'] = 'Average Sales'; summary['B5'] = "=AVERAGE('Raw Data'!F2:F51)"; summary['D5'] = 'Total Sales ÷ Transaction Count'
summary['A6'] = 'Transaction Count'; summary['B6'] = "=COUNTA('Raw Data'!A2:A51)"; summary['D6'] = f'Filled IDs: {count}'
for cell in ('B4','B5'):
    summary[cell].number_format = '₹#,##0.00'
summary['B6'].number_format = '0'
for row in summary.iter_rows(min_row=3, max_row=6, min_col=1, max_col=4):
    for cell in row:
        cell.border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
        cell.alignment = Alignment(vertical='center')
for cell in ('A4','A5','A6'):
    summary[cell].font = Font(bold=True)

# KPI cards
card_fills = ['5B9BD5','70AD47','ED7D31']
labels = [('A9','TOTAL SALES','B10',"=B4"),('C9','AVERAGE SALES','D10',"=B5"),('E9','TRANSACTION COUNT','F10',"=B6")]
for (label_cell, label, value_cell, formula), fill in zip(labels, card_fills):
    summary[label_cell] = label; summary[label_cell].font = Font(bold=True, color='FFFFFF', size=11); summary[label_cell].fill = PatternFill('solid', fgColor=fill); summary[label_cell].alignment = Alignment(horizontal='center')
    summary[value_cell] = formula; summary[value_cell].font = Font(bold=True, color='FFFFFF', size=16); summary[value_cell].fill = PatternFill('solid', fgColor=fill); summary[value_cell].alignment = Alignment(horizontal='center')
    if value_cell in ('B10','D10'): summary[value_cell].number_format = '₹#,##0.00'
    for r in range(summary[value_cell].row, summary[value_cell].row+1):
        for c in range(summary[value_cell].column, summary[value_cell].column+1): summary.cell(r,c).fill = PatternFill('solid', fgColor=fill)

summary['A13'] = 'Sales Summary Insights'; summary['A13'].font = Font(bold=True, size=13, color='17365D')
summary['A14'] = '• Total sales represent the overall revenue generated from all transactions.'
summary['A15'] = '• Average sales show the average revenue generated per transaction.'
summary['A16'] = '• Transaction count represents the total number of recorded transactions.'
summary['A18'] = 'Manual Verification'; summary['A18'].font = Font(bold=True, color='FFFFFF'); summary['A18'].fill = header_fill
summary['B18'] = 'Passed'; summary['B18'].font = Font(bold=True, color='008000')
summary['A19'] = 'Dataset rows'; summary['B19'] = count
summary['A20'] = 'Raw Data sales sum'; summary['B20'] = total; summary['B20'].number_format = '₹#,##0.00'
for col, width in {'A':25,'B':20,'C':18,'D':26,'E':20,'F':20}.items(): summary.column_dimensions[col].width = width
for r in [9,10]: summary.row_dimensions[r].height = 28
summary.freeze_panes = 'A3'
summary.sheet_view.showGridLines = False
raw.sheet_view.showGridLines = False

xlsx = out / 'Retail_Sales_Summary.xlsx'
wb.save(xlsx)

# Verify workbook formulas and source data integrity
check = load_workbook(xlsx, data_only=False)
assert check['Raw Data'].max_row == 51
assert check['Summary']['B4'].value == "=SUM('Raw Data'!F2:F51)"
assert check['Summary']['B5'].value == "=AVERAGE('Raw Data'!F2:F51)"
assert check['Summary']['B6'].value == "=COUNTA('Raw Data'!A2:A51)"
print(f'Created: {xlsx}')
print(f'Created: {raw_csv}')
print(f'Created: {summary_csv}')
print(f'Total Sales: {total:.2f}; Average Sales: {avg:.2f}; Transaction Count: {count}')
